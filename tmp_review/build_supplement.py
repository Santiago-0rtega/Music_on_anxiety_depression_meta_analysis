from pathlib import Path
from io import BytesIO
from docx import Document
from pypdf import PdfReader,PdfWriter,Transformation
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import Paragraph,Table,TableStyle
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.utils import simpleSplit

BASE=Path(r'C:\Users\zannt\OneDrive\Papers sub\Biology Open\Music meta')
ORIG=BASE/'New folder'/'Supplementary material.pdf'
FIG=BASE/'bias_lnRR.pdf'
MAIN=BASE/'Music_meta_analysis_revised.docx'
OUT=Path('review_output/Supplementary_material_BiO_revised.pdf')
W,H=letter
style=ParagraphStyle('cell',fontName='Helvetica',fontSize=9,leading=13,alignment=TA_LEFT)
head=ParagraphStyle('head',fontName='Helvetica-Bold',fontSize=9,leading=13)

def page_base(title,num):
 buf=BytesIO();c=Canvas(buf,pagesize=letter)
 c.setFont('Helvetica-Bold',14);c.drawString(54,742,title)
 c.setFont('Helvetica',9);c.drawCentredString(W/2,36,str(num))
 return buf,c

def finish(buf,c):
 c.showPage();c.save();buf.seek(0);return PdfReader(buf).pages[0]

src=Document(MAIN)
# Figure S3, preserving the supplied vector artwork and placing its legend on the same page.
buf,c=page_base('S6 Additional figure and tables',31)
c.setFont('Helvetica-Bold',11);c.drawString(54,706,'Figure S3. Small-study and time-lag diagnostics')
legend=('Associations between effect sizes and (A) the square root of the inverse effective sample size and '
'(B) publication year. Point size reflects precision; solid lines show fitted meta-regressions, and dashed '
'and dotted lines indicate 95% confidence and prediction intervals, respectively.')
y=224
for line in simpleSplit(legend,'Helvetica',9,504):c.setFont('Helvetica',9);c.drawString(54,y,line);y-=13
figure_page=finish(buf,c)
plot=PdfReader(FIG).pages[0]
scale=0.70
figure_page.merge_transformed_page(plot,Transformation().scale(scale).translate(54,315),over=True)

# Table S4: moved main-paper Table 2, preserving all values.
buf,c=page_base('S6 Additional figure and tables',32)
c.setFont('Helvetica-Bold',11);c.drawString(54,706,'Table S4. Heterogeneity explained by individual moderators')
caption=('Proportion of total heterogeneity (marginal R-squared) explained by individual moderators. '
'Separate univariate meta-regressions included anxiety- and depression-like effect sizes together; '
'estimates are not independent or additive.')
y=682
for line in simpleSplit(caption,'Helvetica',9,504):c.setFont('Helvetica',9);c.drawString(54,y,line);y-=13
rows=[]
for row in src.tables[1].rows:
 rows.append([Paragraph(cell.text.replace('lnRR Marginal','lnRR marginal R²').replace('lnVR Marginal','lnVR marginal R²'),head if not rows else style) for cell in row.cells])
t=Table(rows,colWidths=[244,130,130],repeatRows=1,hAlign='LEFT')
t.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LINEBELOW',(0,0),(-1,0),0.8,colors.black),('LINEBELOW',(0,-1),(-1,-1),0.6,colors.black),('LINEBELOW',(0,1),(-1,-2),0.25,colors.HexColor('#CCCCCC')),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7),('LEFTPADDING',(0,0),(-1,-1),6)]))
tw,th=t.wrapOn(c,504,600);t.drawOn(c,54,y-20-th)
table4=finish(buf,c)

# Table S5: moved main-paper Table 3.
buf,c=page_base('S6 Additional figure and tables',33)
c.setFont('Helvetica-Bold',11);c.drawString(54,706,'Table S5. PECO eligibility framework')
c.setFont('Helvetica',9);c.drawString(54,682,'Population, exposure, comparator and outcome definitions used to scope the review.')
rows=[]
for row in src.tables[2].rows:
 rows.append([Paragraph(cell.text,head if not rows else style) for cell in row.cells])
t=Table(rows,colWidths=[120,384],repeatRows=1,hAlign='LEFT')
t.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('LINEBELOW',(0,0),(-1,0),0.8,colors.black),('LINEBELOW',(0,-1),(-1,-1),0.6,colors.black),('LINEBELOW',(0,1),(-1,-2),0.25,colors.HexColor('#CCCCCC')),('TOPPADDING',(0,0),(-1,-1),9),('BOTTOMPADDING',(0,0),(-1,-1),9),('LEFTPADDING',(0,0),(-1,-1),6)]))
tw,th=t.wrapOn(c,504,600);t.drawOn(c,54,650-th)
table5=finish(buf,c)
writer=PdfWriter()
# Replace the outdated ZIP-archive statement on the original page 19.
buf,c=page_base('S3.1 Supplementary CRIME-Q extraction materials',19)
y=701
paragraphs=[
 'Four Markdown files used to standardize the NotebookLM-assisted CRIME-Q assessment are available in the project repository and the Borealis archive. The repository folder is:',
 'github.com/Santiago-0rtega/Music_on_anxiety_depression_meta_analysis/tree/main/CRIMEQ/notebooklm_workflow/shared_rules',
 '<b>CRIME-Q_CODEBOOK_FOR_NOTEBOOKLM.md.</b> Defines the 20 appraisal items, permitted response options, and decision rules for animal reporting, acoustic interventions, study conduct, risk of bias, reporting quality, and funding or conflict-of-interest considerations.',
 '<b>CRIME-Q_NOTEBOOKLM_CALIBRATION_RULES.md.</b> Gives calibration rules for high-risk or commonly misclassified items, the distinction between behavioral and non-behavioral evidence, and required verbatim supporting quotations.',
 '<b>GENERAL_NOTEBOOKLM_INSTRUCTIONS.md.</b> Describes the extraction workflow, unit of assessment, sections of each paper to inspect, acceptable scores, evidence requirements, and output format.',
 '<b>WIDE_SHEET_COLUMN_TEMPLATE.md.</b> Specifies the columns and ordering of the wide-format extraction table, with one row per study and item-level score, justification, and source evidence.'
]
body=ParagraphStyle('body',fontName='Helvetica',fontSize=10,leading=15,spaceAfter=12)
for para in paragraphs:
 p=Paragraph(para,body);pw,ph=p.wrap(504,700);p.drawOn(c,54,y-ph);y-=ph+14
replacement19=finish(buf,c)
for idx,p in enumerate(PdfReader(ORIG).pages):writer.add_page(replacement19 if idx==18 else p)
for p in (figure_page,table4,table5):writer.add_page(p)
OUT.parent.mkdir(exist_ok=True)
with OUT.open('wb') as f:writer.write(f)
print(OUT.resolve(),len(PdfReader(OUT).pages))
