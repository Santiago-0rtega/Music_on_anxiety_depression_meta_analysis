from copy import deepcopy
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

path='review_output/Music_meta_analysis_BiO_formatted.docx'
d=Document(path)

def insert_like(run, value, anchor, changed):
    node=OxmlElement('w:r')
    if run._r.rPr is not None: node.append(deepcopy(run._r.rPr))
    if changed:
        props=node.find(qn('w:rPr'))
        if props is None:props=OxmlElement('w:rPr');node.insert(0,props)
        mark=props.find(qn('w:highlight'))
        if mark is None:mark=OxmlElement('w:highlight');props.append(mark)
        mark.set(qn('w:val'),'yellow')
    t=OxmlElement('w:t');t.text=value
    if value.startswith(' ') or value.endswith(' '):t.set(qn('xml:space'),'preserve')
    node.append(t);anchor.addnext(node)
    return node

def replace(para, old, new):
    for run in para.runs:
        if old in run.text:
            before,after=run.text.split(old,1)
            run.text=before
            node=insert_like(run,new,run._r,True)
            if after:insert_like(run,after,node,False)
            return
    raise ValueError(f'Phrase not found: {old}')

revisions={
44:[('predominance of negative effect sizes supports a clear directional effect','predominance of negative effect sizes is consistent with an overall directional association'),('music alters central tendency without uniformly reshaping individual variation','music exposure was associated with a shift in central tendency without a consistent change in individual variation')],
46:[('music functions as a sensory form of environmental enrichment','music may act as a sensory form of environmental enrichment'),('Interventions that modulate stress physiology should therefore produce larger effects','Interventions that modulate stress physiology could therefore show larger effects')],
47:[('Timing also influenced','Exposure timing was also associated with'),('this pattern suggests that music primarily shifts baseline affective states rather than enhancing task performance','this pattern is consistent with a possible influence on pretest affective state, although differences in exposure protocols or study design could also contribute'),('prior exposure allows physiological and affective modulation to occur before assessment','prior exposure may permit physiological and affective modulation before assessment')],
49:[('Behavioral assays strongly moderated variability responses','Behavioral assays were associated with marked differences in variability responses'),('music appears to promote convergence toward similar coping responses','the observed pattern is consistent with more similar coping responses'),('Music may amplify pre-existing differences','This pattern could reflect pre-existing differences')],
51:[('Music meta-genre also structured variability','Music meta-genre was also associated with differences in variability'),('they likely capture differences in auditory temporal structure and predictability','they may partly reflect differences in auditory temporal structure and predictability'),('genre-dependent variability likely reflects differences','genre-associated variability could reflect differences')],
52:[('Experimental design also influenced variability','Experimental design was also associated with variability')],
63:[('Music exposure primarily shifts the central tendency of affect-related behavior','The available evidence more consistently supports a shift in average affect-related behavior'),('with changes in behavioral dispersion emerging only under specific experimental and stimulus conditions','while differences in behavioral dispersion appeared in some experimental and stimulus contexts')]
}
for index,items in revisions.items():
    for old,new in items:replace(d.paragraphs[index],old,new)

authors=[
 ('Santiago Ortega','¹*','0000-0002-3518-276X'),
 ('Anna Lenz','¹','0009-0002-5706-7815'),
 ('Erick Lundgren','¹','0000-0001-9893-3324'),
 ('Ayumi Mizuno','¹','0000-0003-0822-5637'),
 ('Sergio Poo Hernandez','¹','0000-0003-0155-645X'),
 ('Shinichi Nakagawa','¹,²†','0000-0002-7765-5182'),
 ('Malgorzata Lagisz','¹,²†','0000-0002-3993-6127')
]
for i,(name,markers,orcid) in enumerate(authors,1):
    p=d.paragraphs[i]
    p.text=''
    r=p.add_run(f'{name}{markers} (ORCID: {orcid})')
    r.font.highlight_color=WD_COLOR_INDEX.YELLOW

anchor=d.paragraphs[8]
details=[
 'Running title: Music and affect in rodents',
 '¹ Collaboration for Open Science and Synthesis in Ecology and Evolution (COSSEE), Department of Biological Sciences, Faculty of Science, University of Alberta, Edmonton, Canada',
 '² School of Biological, Earth and Environmental Sciences, University of New South Wales, Sydney, New South Wales, Australia',
 '* Corresponding author: Santiago Ortega (santiago.ortega@ualberta.ca)',
 '† Co-senior authors: Shinichi Nakagawa and Malgorzata Lagisz'
]
for detail in details:
    p=anchor.insert_paragraph_before()
    p.style=d.styles['Normal']
    p.paragraph_format.space_after=Pt(5)
    r=p.add_run(detail)
    r.font.size=Pt(10)
    r.font.highlight_color=WD_COLOR_INDEX.YELLOW
d.save(path)
print('Revised discussion and title page')
