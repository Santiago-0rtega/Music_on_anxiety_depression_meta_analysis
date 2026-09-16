from pathlib import Path
from docx import Document
from docx.enum.text import WD_LINE_SPACING, WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import re, unicodedata
from difflib import SequenceMatcher
SRC=Path(r'C:\Users\zannt\OneDrive\Papers sub\Biology Open\Music meta\Music_meta_analysis_revised.docx')
BIB=Path(r'C:\Users\zannt\OneDrive\Papers sub\Biology Open\Music meta\New folder\refs.bib')
OUT=Path('review_output');OUT.mkdir(exist_ok=True)

def parse(block):
 m=re.match(r'@(\w+)\{([^,]+),',block)
 if not m:return None
 typ,key=m.groups();s=block[m.end():];d={'type':typ,'key':key};i=0
 while i<len(s):
  m=re.search(r'([A-Za-z_-]+)\s*=\s*',s[i:])
  if not m:break
  name=m.group(1).lower();i+=m.end();q=s[i:i+1]
  if q=='{':
   depth=1;j=i+1
   while j<len(s) and depth:
    if s[j]=='{':depth+=1
    elif s[j]=='}':depth-=1
    j+=1
   val=s[i+1:j-1];i=j
  elif q=='"':
   j=s.find('"',i+1);val=s[i+1:j];i=j+1
  else:
   m2=re.match(r'([^,}\n]+)',s[i:]);val=m2.group(1).strip() if m2 else '';i+=len(m2.group(0)) if m2 else 1
  d[name]=val
 return d

def clean(s):
 accents={"'":'\u0301','"':'\u0308','`':'\u0300','~':'\u0303','^':'\u0302'}
 s=re.sub(r'\\([\'"`~^])\{?([A-Za-z])\}?',lambda m:unicodedata.normalize('NFC',m.group(2)+accents[m.group(1)]),s)
 s=re.sub(r'\\c\{?([A-Za-z])\}?',lambda m:unicodedata.normalize('NFC',m.group(1)+'\u0327'),s)
 s=s.replace('{','').replace('}','').replace('\\&','&').replace('\\_','_').replace('\\%','%')
 s=re.sub(r'\\textit\{([^}]*)\}',r'\1',s)
 s=re.sub(r'\s+',' ',s).strip()
 return s

def initials(given):
 parts=re.split(r'\s+',given.strip());out=[]
 for p in parts:
  if not p:continue
  if re.fullmatch(r'[A-Z]\.?',p):out.append(p[0]+'.')
  elif '-' in p:out.append('-'.join(x[0].upper()+'.' for x in p.split('-') if x))
  else:out.append(p[0].upper()+'.')
 return ' '.join(out)

def authors(s):
 arr=[]
 for a in re.split(r'\s+and\s+',clean(s)):
  if ',' in a:
   sur,given=a.split(',',1);arr.append(sur.strip()+', '+initials(given))
  else:
   words=a.split();arr.append((words[-1]+', '+initials(' '.join(words[:-1]))) if len(words)>1 else a)
 if len(arr)>1:return ', '.join(arr[:-1])+' and '+arr[-1]
 return arr[0] if arr else ''

def norm(s):return re.sub(r'[^a-z0-9]+','',unicodedata.normalize('NFKD',s).lower())
entries=[parse(x) for x in re.split(r'(?m)(?=^@\w+\{)',BIB.read_text(encoding='utf-8-sig')) if x.startswith('@')]
entries=[x for x in entries if x]
d=Document(SRC)
matched=0;unmatched=[];ambig=[]
for p in d.paragraphs:
 if p.style.name!='Bibliography':continue
 old=p.text
 if old.startswith('Posit team. 2026.'):
  p.clear();p.add_run('Posit team (2026). RStudio: Integrated Development Environment for R. Posit Software, PBC. http://www.posit.co/.');matched+=1;continue
 if old.startswith('R Core Team. 2025.'):
  p.clear();p.add_run('R Core Team (2025). R: A Language and Environment for Statistical Computing. R Foundation for Statistical Computing. https://www.R-project.org/.');matched+=1;continue
 year=re.search(r'\b(19|20)\d{2}\b',old)
 first=norm(old.split(',')[0])
 candidates=[e for e in entries if norm(e.get('author','').split(',')[0])==first and e.get('year')==(year.group(0) if year else '')]
 if not candidates and year:
  candidates=[e for e in entries if norm(e['key']).startswith(first) and year.group(0) in e['key']]
 if not candidates and first=='witte' and year:
  candidates=[e for e in entries if 'witte' in norm(e['key']) and year.group(0) in e['key']]
 if len(candidates)>1:
  titlepart=re.search(r'[“\"]([^”\"]+)',old)
  if titlepart:candidates=sorted(candidates,key=lambda e:SequenceMatcher(None,norm(e.get('title','')),norm(titlepart.group(1))).ratio(),reverse=True)[:1]
 if len(candidates)!=1:
  unmatched.append(old[:100]);continue
 e=candidates[0];matched+=1
 if e['key'].startswith('Airaksinen2024'):
  e=dict(e);e['author']='Spytska, Liana'
 if e['key'].startswith('Wickham2019'):
  e=dict(e);e['author']=e.get('author','').replace("D', Lucy and Mcgowan, Agostino","D'Agostino McGowan, Lucy")
 au=authors(e.get('author',''));yr=e.get('year','');title=clean(e.get('title','')).rstrip('. ')
 typ=e['type'].lower();j=clean(e.get('journal','') or e.get('booktitle',''));vol=clean(e.get('volume',''));pages=clean(e.get('pages','')).replace('--','-').replace('–','-');doi=clean(e.get('doi',''))
 if typ=='article':
  new=f'{au} ({yr}). {title}. {j}'
  if vol:new+=f' {vol}'
  if pages:new+=f', {pages.rstrip("-")}'
  new+='.'
 elif typ in ('book','manual'):
  new=f'{au} ({yr}). {title}. '+clean(e.get('publisher',''))+'.'
 elif typ in ('incollection','inbook'):
  new=f'{au} ({yr}). {title}. In {j}'
  if pages:new+=f', pp. {pages}'
  new+='.'
 else:new=f'{au} ({yr}). {title}. {j}.'
 # BiO supplies matched DOIs at proof stage; source BibTeX has several malformed DOIs.
 p.clear();p.add_run(new)
# Style formatting while retaining all substantive body content, equations, tables, and figures.
for p in d.paragraphs:
 p.paragraph_format.line_spacing=1.5
 if p.style.name=='Bibliography':
  p.paragraph_format.left_indent=Pt(18);p.paragraph_format.first_line_indent=Pt(-18);p.paragraph_format.space_after=Pt(6)
# User-approved clarification, shown in yellow because it is outside references.
workflow_url='https://github.com/Santiago-0rtega/Music_on_anxiety_depression_meta_analysis/tree/main/CRIMEQ/notebooklm_workflow/shared_rules'
additions={
 127: ' The CRIME-Q workflow files are also included in the Borealis archive. The codebook defines the appraisal items and scoring rules; the calibration rules guide uncertain ratings; the general instructions describe source-grounded checks; and the wide-sheet template defines the assessment columns. The files can be accessed directly at '+workflow_url+'.',
 137: ' The CRIME-Q workflow files are included in both the project repository and the Borealis archive.'
}
for index,addition in additions.items():
 run=d.paragraphs[index].add_run(addition)
 run.font.highlight_color=7
# Author-approved abstract shortening and citation correction.
abstract=('Music exposure may influence affect-related behavior in rodents, but results vary across studies, and effects on behavioral variability remain unclear. '
 'We conducted a preregistered systematic review and multilevel meta-analysis of 20 experimental studies (298 effect sizes) testing music exposure in laboratory rodents. '
 'Using the log response ratio (lnRR) and log variability ratio (lnVR), we estimated changes in average anxiety- and depression-like behavior and behavioral dispersion while accounting for non-independence among effects. '
 'We also assessed reporting, methodological quality, and risk of bias using CRIME-Q. '
 'Music exposure was associated with an approximately 20% reduction in anxiety- and depression-like behavior relative to control conditions, with no overall change in behavioral variability. '
 'Exploratory moderator analyses suggested context-dependent dispersion: anxiety-like outcomes tended toward greater variability and depression-like outcomes toward lower variability, while behavioral assay and music meta-genre explained comparatively more heterogeneity in lnVR. '
 'Reporting of sample-size justification, allocation, blinding, attrition, and acoustic delivery was often incomplete. '
 'The average reduction applies primarily to the studied laboratory contexts; variability effects and generalization beyond predominantly young adult rodents and short-term exposure require caution.')
ap=d.paragraphs[10];ap.clear();ar=ap.add_run(abstract);ar.font.highlight_color=7
for p in d.paragraphs[:155]:
 for r in p.runs:
  if 'Airaksinen et al. 2024' in r.text:
   r.text=r.text.replace('Airaksinen et al. 2024','Spytska 2024')
   r.font.highlight_color=7
# Repair stale cross-references created when former Table S1 moved into the main text.
def retarget_hyperlink(paragraph_index,old_anchor,new_anchor,new_text):
 for h in d.paragraphs[paragraph_index]._p.xpath('./w:hyperlink'):
  if h.get(qn('w:anchor'))==old_anchor:
   h.set(qn('w:anchor'),new_anchor)
   for t in h.xpath('.//w:t'):
    t.text=new_text
   for r in h.xpath('.//w:r'):
    rpr=r.find(qn('w:rPr'))
    if rpr is None:rpr=OxmlElement('w:rPr');r.insert(0,rpr)
    hi=rpr.find(qn('w:highlight'))
    if hi is None:hi=OxmlElement('w:highlight');rpr.append(hi)
    hi.set(qn('w:val'),'yellow')
for idx in (32,34,45):retarget_hyperlink(idx,'tab:moderators','tab:heterogeneity','2')
retarget_hyperlink(67,'tab:peco','tab:peco','3')
retarget_hyperlink(57,'fig','fig:crimeq','8')
caption=d.paragraphs[38]._p
bookmark_start=OxmlElement('w:bookmarkStart');bookmark_start.set(qn('w:id'),'9001');bookmark_start.set(qn('w:name'),'tab:heterogeneity')
bookmark_end=OxmlElement('w:bookmarkEnd');bookmark_end.set(qn('w:id'),'9001')
caption.insert(1,bookmark_start);caption.append(bookmark_end)
for run in d.paragraphs[78].runs:
 if 'Table S1' in run.text:run.text=run.text.replace('Table S1','Table 1');run.font.highlight_color=7
for run in d.paragraphs[142].runs:
 if run.text=='60 days).':run.text='>60 days).';run.font.highlight_color=7
# Move the three author-selected display items to the supplement.
def unlink_to_supplement(paragraph_index,anchor,display):
 p=d.paragraphs[paragraph_index]._p
 for h in list(p.xpath('./w:hyperlink')):
  if h.get(qn('w:anchor'))!=anchor:continue
  for t in h.xpath('.//w:t'):t.text=display
  for r in h.xpath('.//w:r'):
   rpr=r.find(qn('w:rPr'))
   if rpr is None:rpr=OxmlElement('w:rPr');r.insert(0,rpr)
   hi=rpr.find(qn('w:highlight'))
   if hi is None:hi=OxmlElement('w:highlight');rpr.append(hi)
   hi.set(qn('w:val'),'yellow')
   h.addprevious(r)
  p.remove(h)
for idx in (32,34,45):unlink_to_supplement(idx,'tab:heterogeneity','S4')
for idx in (67,72):unlink_to_supplement(idx,'tab:peco','S5')
unlink_to_supplement(40,'fig:bias','S3')
for idx in (42,57):retarget_hyperlink(idx,'fig:crimeq','fig:crimeq','7')
for r in d.paragraphs[40].runs:
 if 'Fig. S2' in r.text or 'Fig S3' in r.text:
  r.text=r.text.replace('Fig. S2','Fig. S1').replace('Fig S3','Fig. S2')
  r.font.highlight_color=7
for r in d.paragraphs[154].runs:
 if 'Figure 8.' in r.text:r.text=r.text.replace('Figure 8.','Figure 7.');r.font.highlight_color=7
for t in (d.tables[1]._element,d.tables[2]._element):t.getparent().remove(t)
for p in [d.paragraphs[idx]._p for idx in (38,69,151,152)]:p.getparent().remove(p)
# Keep the reference list in alphabetical order after correcting the misattributed article.
refs=[p for p in d.paragraphs if p.style.name=='Bibliography']
body=d.element.body
for p in refs:body.remove(p._p)
for p in sorted(refs,key=lambda x:norm(x.text.split(',')[0])):
 body.insert(len(body)-1,p._p)
for t in d.tables:
 for row in t.rows:
  for c in row.cells:
   for p in c.paragraphs:p.paragraph_format.line_spacing=1.5
for s in d.sections:
 ln=s._sectPr.find(qn('w:lnNumType'))
 if ln is None:ln=OxmlElement('w:lnNumType');s._sectPr.append(ln)
 ln.set(qn('w:countBy'),'1');ln.set(qn('w:restart'),'continuous')
 footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
 for child in list(footer._p):footer._p.remove(child)
 r=footer.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
output=OUT/'Music_meta_analysis_BiO_formatted.docx';d.save(output)
(OUT/'reference_matching_log.txt').write_text(f'Matched {matched} of 97 bibliography entries to supplied BibTeX.\nUnmatched original entries retained:\n'+'\n'.join(unmatched),encoding='utf-8')
print(output.resolve());print('matched',matched,'unmatched',len(unmatched))
