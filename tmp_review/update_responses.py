from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
SRC=Path(r'C:\Users\zannt\OneDrive\Papers sub\Biology Open\Music meta\Comments from the Reviewers.docx')
OUT=Path('review_output/Reviewer_responses_lines_updated.docx')
d=Document(SRC)
changes={
 9:[('lines 26--29','lines 28–31'),('lines 715--722','lines 695–704')],
 12:[('lines 80--82','lines 79–83')],
 16:[('lines 94--97','lines 96–98')],
 24:[('lines 130--133','lines 132–136')],
 26:[('lines 710--712','lines 690–694'),('lines 137--138','lines 140–142'),('lines 551--554','lines 525–528')],
 28:[('lines 203--207','lines 178–183')],
 30:[('(line 153)','(Table 1)')],
 32:[('lines 319--329','lines 313–326')],
 37:[('lines 184--189 and 213--218','lines 185–190 and 214–219'),('lines 184--189','lines 185–190 and 214–219')],
 43:[('lines 667--670','lines 672–676')],
 45:[('lines 323--329','lines 319–326')],
 47:[('We have added human','We have added human evidence'),('lines 332--336','lines 328–333'),('lines 454--457','lines 454–457')],
 49:[('lines 597-603','lines 595–601'),('lines 277--283','lines 273–279')],
 50:[('lines 434--448','lines 432–448')],
 52:[('lines 359--360','lines 330–333')],
 59:[('lines 620–626','lines 602–608')],
 61:[('lines 194–196','lines 170–172')],
 70:[('lines 109–120','lines 111–122'),('lines 148–180','lines 151–155 and Table 1')],
 76:[('lines 130–135','lines 132–136')],
 77:[('lines 301–307','lines 273–279')],
 85:[('lines 173–180','Table 1'),('lines 155–158','Table 1')],
 88:[('lines 41–42','lines 42–48'),('Airaksinen et al., 2024','Spytska, 2024')],
 90:[('lines 139–142','lines 142–145')],
 106:[('Table 1','Table S4')],
 115:[('lines 150–180','lines 151–155 and Table 1')]
}

def insert_run_like(run,text,after,highlight):
 el=OxmlElement('w:r')
 if run._r.rPr is not None:el.append(deepcopy(run._r.rPr))
 if highlight:
  rpr=el.find(qn('w:rPr'))
  if rpr is None:rpr=OxmlElement('w:rPr');el.insert(0,rpr)
  hi=rpr.find(qn('w:highlight'))
  if hi is None:hi=OxmlElement('w:highlight');rpr.append(hi)
  hi.set(qn('w:val'),'yellow')
 t=OxmlElement('w:t');t.text=text
 if text.startswith(' ') or text.endswith(' '):t.set(qn('xml:space'),'preserve')
 el.append(t);after.addnext(el)
 return el

def change_text(p,old,new):
 for run in p.runs:
  if old in run.text:
   before,after=run.text.split(old,1);run.text=before
   changed=insert_run_like(run,new,run._r,True)
   if after:insert_run_like(run,after,changed,False)
   return True
 return False

for i,items in changes.items():
 for old,new in items:
  if not change_text(d.paragraphs[i],old,new):
   if old in d.paragraphs[i].text:
    p=d.paragraphs[i];p.text=p.text.replace(old,new,1)
    if p.runs:p.runs[0].font.highlight_color=7
    print('FULL-PARAGRAPH HIGHLIGHT',i,repr(old))
# Correct claims about the pending DOI and describe the requested direct CRIME-Q access.
replacements={
 66:'We thank the reviewer for this suggestion. The data, analysis code, metadata, and supporting materials are available in the project repository and the Borealis archive. The manuscript retains a DOI placeholder, which will be replaced with the permanent Borealis DOI before submission. We will confirm the archive’s reuse licence in the final data-availability statement.',
 122:'We thank the reviewer for this suggestion. We have included the adapted CRIME-Q codebook, with item definitions, response options, and operational decision rules, in the Supplementary Material. The other workflow files are in both the project repository and the Borealis archive. The manuscript now links directly to the repository folder and describes the calibration rules, general instructions, and wide-sheet template (lines 785–796).'
}
for i,new in replacements.items():
 p=d.paragraphs[i]
 for r in p.runs:r.text=''
 r=p.add_run(new);r.font.highlight_color=7
OUT.parent.mkdir(exist_ok=True);d.save(OUT)
print(OUT.resolve())
